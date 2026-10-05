"""Gera DT_na_Pratica_Slides.pptx com a identidade HayaiDataSystems.

Vídeos: cada módulo tem um botão "Assistir vídeo" que abre um slide de vídeo
OCULTO (pulado na apresentação normal). O slide oculto tem um botão "Voltar".

Animações (anim.py): cada slide tem uma coreografia de entrada que roda sozinha.
Só pedem clique a revelação da resposta na correção da prova e a pergunta do estudo de caso.
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

from brand import *  # noqa: F401,F403
import brand as B
import anim as A
from anim import beat
from content import MODULES, SCHEDULE, CHALLENGE, EXAM, VIDEOS

W, H = 13.333, 7.5
FONT = "Segoe UI"
FONT_B = "Segoe UI Semibold"
FONT_L = "Segoe UI Light"
STAGES = ["Empatizar", "Definir", "Idear", "Prototipar", "Testar"]
STAGE_TXT = ["Entender as pessoas e o contexto", "Encontrar o problema certo", "Gerar muitas soluções", "Tornar a ideia tangível", "Validar com pessoas reais"]
MODULE_STAGE = {1: None, 2: "Empatizar", 3: "Definir", 4: "Idear", 5: "Prototipar", 6: None}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def C(h):
    return RGBColor.from_string(h)


def set_alpha(shape, alpha_pct):
    """Transparência do preenchimento sólido (alpha 0-100 = opacidade)."""
    clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha_pct * 1000)))


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, lw=1.0, alpha=None, dash=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = C(fill)
        if alpha is not None:
            set_alpha(s, alpha)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = C(line)
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.text = ""
    return A.track(s)


def text(slide, x, y, w, h, content, size=16, color=TEXT, bold=False, font=None, align="l",
         anchor="t", italic=False, spacing=None, line_spacing=None, shape=None):
    """content: str | list[str] (parágrafos) | list[list[(txt, {opts})]] (runs)."""
    if shape is None:
        tb = A.track(slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)))
    else:
        tb = shape
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = content if isinstance(content, list) else [content]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if line_spacing:
            para.line_spacing = line_spacing
        if spacing is not None and i < len(paras) - 1:
            para.space_after = Pt(spacing)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.name = o.get("font", font or (FONT_B if o.get("bold", bold) else FONT))
            f.color.rgb = C(o.get("color", color))
            if o.get("cs"):
                r._r.get_or_add_rPr().set("spc", str(o["cs"]))
    return tb


def bullets(slide, x, y, w, h, items, size=17, color=TEXT, dot=LIME, spacing=12):
    tb = A.track(slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.32))))
        pPr.set("indent", str(-Emu(Inches(0.32))))
        bc = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(bc, qn("a:srgbClr")).set("val", dot)
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "●")
        if i < len(items) - 1:
            p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = it
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = C(color)
    return tb


def arrow_line(slide, x1, y1, x2, y2, color=MUTED, w=1.5, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = C(color)
    c.line.width = Pt(w)
    if head:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")
    return A.track(c)


def circle(slide, cx, cy, d, fill, label=None, size=14, color=BG, bold=True, alpha=None):
    s = box(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, shape=MSO_SHAPE.OVAL, alpha=alpha)
    if label is not None:
        text(slide, 0, 0, 0, 0, label, size=size, color=color, bold=bold, align="c", anchor="m", shape=s)
        s.text_frame.margin_left = s.text_frame.margin_right = 0
    return s


def photo_fill(slide, shape, path, ring=LIME):
    """Preenche uma forma (ex.: círculo) com uma foto, mantendo a forma e suas animações."""
    _, rid = slide.part.get_or_add_image_part(str(path))
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill"):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    blip = etree.SubElement(spPr, qn("a:blipFill"))
    etree.SubElement(blip, qn("a:blip")).set(qn("r:embed"), rid)
    etree.SubElement(etree.SubElement(blip, qn("a:stretch")), qn("a:fillRect"))
    spPr.remove(blip)
    spPr.insert(list(spPr).index(spPr.find(qn("a:prstGeom"))) + 1, blip)
    shape.line.color.rgb = C(ring)
    shape.line.width = Pt(2.25)
    return shape


def new_slide(notes=None, transition="fade"):
    A.finish()
    s = prs.slides.add_slide(BLANK)
    A.begin(s, transition)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = C(BG)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def chrome(s, tag=None, tag_color=LIME):
    if tag:
        text(s, 0.6, 0.32, 9, 0.3, [[(tag.upper(), {"cs": 200})]], size=11, color=tag_color, bold=True)
    text(s, 0.6, 7.0, 8, 0.3, f"{B.BRAND}  ·  {B.COURSE}", size=10, color=DIM)
    s.shapes.add_picture(str(ASSETS / "logo_icon_t.png"), Inches(W - 1.12), Inches(6.88), height=Inches(0.42))


def title(s, t, y=0.62, size=32, w=12.1):
    with beat("fade", dur=0.4):
        return text(s, 0.6, y, w, 0.8, t, size=size, bold=True, color=TEXT)


def circuit(s, x, y, color=BLUE, scale=1.0, alpha=None):
    """Motivo de circuito inspirado no logo: linhas com nós."""
    k = scale
    paths = [(0, 0, 1.6, 0), (0, 0.45, 1.1, 0.45), (1.1, 0.45, 1.45, 0.8), (1.45, 0.8, 2.3, 0.8), (0, 0.9, 0.6, 0.9), (0.6, 0.9, 0.9, 1.25), (0.9, 1.25, 1.9, 1.25)]
    # o "sinal" percorre as trilhas e acende os nós no fim
    with beat("wipe", dur=0.9):
        for x1, y1, x2, y2 in paths:
            c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + x1 * k), Inches(y + y1 * k), Inches(x + x2 * k), Inches(y + y2 * k))
            c.line.color.rgb = C(color)
            c.line.width = Pt(2.2 * k)
            A.track(c)
    with beat("zoom", gap=0.75, stagger=0.12):
        for nx, ny in [(1.6, 0), (2.3, 0.8), (1.9, 1.25)]:
            d = 0.22 * k
            box(s, x + nx * k, y + ny * k - d / 2, d, d, fill=BG, line=color, shape=MSO_SHAPE.OVAL, lw=2.2 * k)


def button(s, x, y, w, h, label, fill=LIME, fg=BG, size=14):
    b = box(s, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, 0, 0, 0, 0, label, size=size, color=fg, bold=True, align="c", anchor="m", shape=b)
    return b


def stage_bar(s, active, x=0.6, y=6.35, w=7.0):
    """Mini-barra das 5 etapas, destacando a atual."""
    gap = 0.08
    pw = (w - gap * 4) / 5
    for i, st in enumerate(STAGES):
        col = STAGE_COLORS[st]
        on = st == active
        p = box(s, x + i * (pw + gap), y, pw, 0.34, fill=col if on else SURFACE, line=None if on else LINE,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        text(s, 0, 0, 0, 0, st, size=10, color=BG if on else MUTED, bold=on, align="c", anchor="m", shape=p)


# ---------------------------------------------------------------- slide types
def s_cover():
    s = new_slide("Boas-vindas. Apresente-se em 1 minuto e explique o formato: 3 encontros de 4 horas, muita prática e um desafio real da empresa no final.",
                  transition="black")
    with beat("zoom", dur=0.8):
        A.track(s.shapes.add_picture(str(ASSETS / "logo_icon_t.png"), Inches(7.75), Inches(1.0), height=Inches(3.4)))
    with beat("fade", gap=0.45):
        text(s, 0.8, 1.2, 7, 0.4, [[("TREINAMENTO IN-COMPANY", {"cs": 300})]], size=13, color=LIME, bold=True)
    with beat("rise", gap=0.2):
        text(s, 0.8, 1.75, 7.4, 2.2, [[("Design Thinking", {})], [("na Prática", {"color": LIME})]], size=54, bold=True, line_spacing=0.95)
    with beat("rise", gap=0.25):
        text(s, 0.8, 3.9, 7, 0.6, B.SUBTITLE, size=20, color=MUTED)
    with beat("fade", gap=0.25):
        text(s, 0.8, 5.05, 7, 0.9, [[("Instrutor  ", {"color": DIM}), (B.INSTRUCTOR, {"bold": True})], [(B.HOURS, {"color": MUTED})]], size=15, spacing=6)
    circuit(s, 9.3, 5.2, color=BLUE, scale=1.4)
    text(s, 0.8, 6.85, 8, 0.3, f"{B.BRAND}  ·  {B.TAGLINE}", size=11, color=DIM)
    return s


def s_welcome():
    s = new_slide("Mostre o vídeo de boas-vindas (botão) ou apresente ao vivo. Deixe claro o que cada pessoa leva para casa.")
    chrome(s, "Boas-vindas")
    title(s, "O que você vai levar deste treinamento")
    items = [("Um método", "Um passo a passo para resolver problemas complexos com foco no cliente."),
             ("Ferramentas prontas", "Matriz CSD, persona, mapa de empatia, Crazy 8s, matriz de priorização e mais."),
             ("Um projeto real", "Sua equipe aplica tudo em um desafio da própria empresa.")]
    for i, (h, b) in enumerate(items):
        x = 0.6 + i * 4.1
        with beat("rise", gap=0.2):
            box(s, x, 1.85, 3.8, 3.3, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, x + 0.35, 3.15, 3.2, 0.5, h, size=20, bold=True)
            text(s, x + 0.35, 3.7, 3.2, 1.4, b, size=14, color=MUTED)
        with beat("zoom", gap=0.15):
            circle(s, x + 0.65, 2.55, 0.7, [LIME, TEAL, BLUE][i], str(i + 1), size=18)
    return s


def s_about():
    s = new_slide("Personalize este slide com sua trajetória. Fale de projetos reais em que você aplicou Design Thinking ou análise de dados. Credibilidade vende o treinamento.")
    chrome(s, "Quem conduz")
    title(s, "Sobre o instrutor")
    foto = ASSETS / "instrutor.png"  # foto quadrada; o círculo faz o recorte
    with beat("zoom", dur=0.6):
        if foto.exists():
            photo_fill(s, circle(s, 2.3, 3.7, 2.6, SURFACE2), foto)
        else:
            circle(s, 2.3, 3.7, 2.6, SURFACE2, "Sua foto", size=14, color=MUTED, bold=False)
    with beat("rise", gap=0.25):
        text(s, 4.3, 1.9, 8, 0.6, B.INSTRUCTOR, size=28, bold=True)
        text(s, 4.3, 2.55, 8, 0.4, f"Fundador da {B.BRAND}", size=16, color=LIME)
    with beat("rise", gap=0.3):
        bullets(s, 4.3, 3.3, 8.2, 3, [
        "Tecnólogo em Análise e Desenvolvimento de Sistemas (FIAP)",
        "MBAs em Gestão da Qualidade (Anhanguera) e Project Management (UNISAL)",
        "Bacharel em Administração (Anhanguera)",
        "Certificação Design Thinking – Process (40h) e cursos de IA, Python e Dados",
        "8 anos na indústria no Japão · hoje, dados e automação (Big Data)",
        ], size=16, dot=TEAL)
    return s


def s_rules():
    s = new_slide("Combine as regras com a turma e peça que acrescentem uma. Regras combinadas são mais respeitadas que regras impostas.")
    chrome(s, "Combinados")
    title(s, "Nossos combinados")
    rules = [("Celular no silencioso", "Usaremos o celular em algumas dinâmicas. Fora delas, foco total."),
             ("Sem julgamento", "Toda ideia é bem-vinda. Críticas só na hora certa."),
             ("Mão na massa", "Aprendemos fazendo: participe das dinâmicas."),
             ("Pontualidade", "Começamos e voltamos dos intervalos no horário.")]
    for i, (h, b) in enumerate(rules):
        x = 0.6 + (i % 2) * 6.15
        y = 1.8 + (i // 2) * 2.3
        with beat("rise", gap=0.18):
            box(s, x, y, 5.9, 2.0, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
            text(s, x + 1.1, y + 0.38, 4.6, 0.5, h, size=19, bold=True)
            text(s, x + 1.1, y + 0.9, 4.6, 1.0, b, size=14, color=MUTED)
        with beat("zoom", gap=0.15):
            circle(s, x + 0.6, y + 0.62, 0.55, [LIME, GREEN, TEAL, BLUE][i], "✓", size=16)
    return s


def s_agenda():
    s = new_slide("Mostre a jornada completa. Destaque que o Desafio Final usa um problema real da empresa, que começa a ser trabalhado já no Encontro 1.")
    chrome(s, "Agenda")
    title(s, "Nossa jornada: 12 horas em 3 encontros")
    mods = {1: [], 2: [], 3: []}
    for m in MODULES:
        mods[m["encontro"]].append(m)
    for i, enc in enumerate(SCHEDULE):
        x = 0.6 + i * 4.1
        with beat("rise", gap=0.35 if i else 0.15):
            box(s, x, 1.75, 3.8, 4.75, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
            text(s, x + 0.3, 1.95, 3.3, 0.35, [[(f"ENCONTRO {i + 1}  ·  4H", {"cs": 150})]], size=11, color=[LIME, TEAL, BLUE][i], bold=True)
            text(s, x + 0.3, 2.3, 3.3, 0.9, enc["title"].split("— ")[1], size=19, bold=True)
        yy = 3.3
        extras = {1: ["Dinâmicas 1 e 2"], 2: ["Dinâmicas 3 e 4"], 3: ["Desafio Final", "Prova"]}[i + 1]
        items = [(MODULE_COLORS[m["n"]], f"{m['n']:02d}", m["title"], BG) for m in mods[i + 1]] + [(SURFACE2, "★", e, LIME) for e in extras]
        for fill, lab, txt, fg in items:
            with beat("fade", gap=0.1):
                circle(s, x + 0.55, yy + 0.25, 0.46, fill, lab, size=12, color=fg)
                text(s, x + 0.95, yy + 0.02, 2.7, 0.5, txt, size=14, bold=True, anchor="m")
            yy += 0.65
    return s


def s_section(m):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(f"Abertura do Módulo {m['n']}. Leia os objetivos em voz alta. Se for exibir o vídeo agora, clique em 'Assistir vídeo'; caso contrário, apenas avance: o slide do vídeo fica oculto e é pulado.",
                  transition="black")
    with beat("zoom", dur=0.7):
        text(s, 0.45, 0.95, 5, 2.0, f"{m['n']:02d}", size=120, bold=True, color=col, font=FONT_B)
    with beat("fade", gap=0.3):
        text(s, 0.6, 0.55, 5, 0.4, [[(f"ENCONTRO {m['encontro']}  ·  MÓDULO {m['n']}", {"cs": 250})]], size=12, color=col, bold=True)
    with beat("rise", gap=0.15):
        text(s, 0.6, 2.55, 7.4, 1.4, m["title"], size=38, bold=True, anchor="b", line_spacing=0.95)
    with beat("rise", gap=0.2):
        text(s, 0.6, 4.02, 7.4, 0.5, m["short"], size=18, color=MUTED)
    with beat("rise", gap=0.3):
        box(s, 8.3, 2.3, 4.45, 3.75, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
        text(s, 8.6, 2.5, 4, 0.4, [[("OBJETIVOS", {"cs": 200})]], size=11, color=col, bold=True)
    with beat("fade", gap=0.3):
        bullets(s, 8.6, 2.9, 3.95, 3.1, m["objectives"], size=13.5, dot=col, spacing=8)
    stage = MODULE_STAGE[m["n"]]
    if stage:
        with beat("wipe", gap=0.2):
            stage_bar(s, stage, y=5.45)
    with beat("zoom", gap=0.25):
        btn = button(s, 0.6, 4.7, 3.3, 0.55, "▶  Assistir vídeo do módulo", fill=col)
    circuit(s, 10.3, 0.7, color=col, scale=1.1)
    text(s, 0.6, 7.0, 8, 0.3, f"{B.BRAND}  ·  {B.COURSE}", size=10, color=DIM)
    return s, btn


def s_video(key, label, back_slide, col=LIME):
    s = new_slide("SLIDE OCULTO: só aparece quando o botão 'Assistir vídeo' é clicado. Para inserir seu vídeo: Inserir > Vídeo > Este Dispositivo, e redimensione sobre o quadro tracejado. Depois de inserido, você pode apagar o texto de instrução.",
                  transition="black")
    s._element.set("show", "0")
    chrome(s, f"Vídeo  ·  {label}", tag_color=col)
    text(s, 0.6, 0.66, 9.6, 0.6, VIDEOS[key], size=21, bold=True)
    fw, fh = 9.6, 5.4
    fx, fy = (W - fw) / 2, 1.45
    box(s, fx, fy, fw, fh, fill=SURFACE, line=col, lw=2, dash=True)
    circle(s, W / 2, fy + 1.9, 1.3, col, "▶", size=34)
    text(s, fx + 0.5, fy + 2.85, fw - 1, 0.5, "Espaço para o seu vídeo", size=22, bold=True, align="c")
    text(s, fx + 0.5, fy + 3.45, fw - 1, 1.2, [
        "No PowerPoint: Inserir  >  Vídeo  >  Este Dispositivo…",
        "e ajuste o vídeo sobre este quadro (16:9).",
        [(f"Arquivo sugerido: videos/{key}.mp4", {"color": col})],
    ], size=14, color=MUTED, align="c", spacing=4)
    back = button(s, W - 3.0, 0.62, 2.4, 0.55, "↩  Voltar", fill=SURFACE2, fg=TEXT)
    back.click_action.target_slide = back_slide
    return s


def s_quote(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    with beat("zoom", dur=0.6):
        text(s, 1.0, 1.3, 2, 2, "“", size=160, color=col, font="Georgia", bold=True)
    with beat("words", gap=0.35):
        text(s, 1.6, 2.6, 10.2, 2.4, sp["quote"], size=40, bold=True, line_spacing=1.05)
    with beat("rise", gap=0.3, after=True):
        text(s, 1.6, 5.1, 10, 0.5, "— " + sp["author"], size=16, color=MUTED)
    return s


def s_cards(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    with beat("fade", gap=0.2):
        text(s, 0.6, 1.45, 12, 0.5, sp["lead"], size=17, color=MUTED)
    n = len(sp["cards"])
    cw = (12.13 - 0.3 * (n - 1)) / n
    for i, (h, b) in enumerate(sp["cards"]):
        x = 0.6 + i * (cw + 0.3)
        with beat("rise", gap=0.18 if i else 0.25):
            box(s, x, 2.35, cw, 3.9, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
            text(s, x + 0.35, 3.55, cw - 0.7, 0.9, h, size=19, bold=True)
            text(s, x + 0.35, 4.55, cw - 0.7, 1.65, b, size=14.5, color=MUTED)
        with beat("zoom", gap=0.15):
            circle(s, x + 0.65, 3.0, 0.62, col, str(i + 1), size=17)
    return s


def s_split(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    for i, (ht, items) in enumerate([(sp["left_title"], sp["left"]), (sp["right_title"], sp["right"])]):
        x = 0.6 + i * 6.25
        hi = i == 1
        with beat("rise", gap=0.6 if hi else 0.2):
            box(s, x, 1.7, 5.9, 4.75, fill=SURFACE if not hi else SURFACE2, line=col if hi else None, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
            text(s, x + 0.4, 1.95, 5.1, 0.6, ht, size=22, bold=True, color=col if hi else MUTED)
        with beat("fade", gap=0.25):
            bullets(s, x + 0.4, 2.75, 5.2, 3.5, items, size=16, dot=col if hi else DIM, spacing=14)
    with beat("zoom", gap=0.4):
        circle(s, 6.675, 4.05, 0.6, BG, "vs", size=14, color=MUTED)
    return s


def s_bullets(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    with beat("rise", gap=0.25):
        bullets(s, 0.6, 1.8, 7.3, 4.6, sp["bullets"], size=18, dot=col, spacing=16)
    with beat("rise", gap=0.5):
        box(s, 8.4, 1.8, 4.33, 4.4, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    with beat("zoom", gap=0.3):
        text(s, 8.7, 1.85, 1.5, 1.2, "“", size=80, color=col, font="Georgia", bold=True)
    with beat("words", gap=0.2):
        text(s, 8.75, 3.0, 3.7, 3.0, sp["highlight"], size=21, bold=True, line_spacing=1.1)
    return s


def s_stat(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    cols = [LIME, TEAL, BLUE]
    for i, (num, lab) in enumerate(sp["stats"]):
        x = 0.6 + i * 4.1
        with beat("rise", gap=0.35 if i else 0.25):
            box(s, x, 1.9, 3.8, 3.2, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, x + 0.3, 3.8, 3.2, 1.0, lab, size=16, color=MUTED, align="c")
        with beat("zoom", gap=0.15, dur=0.55):
            text(s, x, 2.25, 3.8, 1.4, num, size=66, bold=True, color=cols[i], align="c")
        if i < 2:
            with beat("wipe", gap=0.25, dur=0.3):
                arrow_line(s, x + 3.82, 3.5, x + 4.08, 3.5, color=DIM)
    with beat("fade", gap=0.4):
        text(s, 0.6, 5.45, 12, 0.8, sp["caption"], size=15, color=MUTED, italic=True)
    return s


def s_pillars(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    d = 2.7
    centers = [(3.0, 3.35), (4.55, 3.35), (3.775, 4.7)]
    cols = [LIME, TEAL, BLUE]
    labels = [(2.35, 2.85), (5.2, 2.85), (3.775, 5.35)]
    for i, ((cx, cy), c, (lx, ly), (h, q, tag)) in enumerate(zip(centers, cols, labels, sp["items"])):
        with beat("zoom", gap=0.3 if i else 0.2, dur=0.6):
            circle(s, cx, cy, d, c, alpha=38)
        with beat("fade", gap=0.2):
            text(s, lx - 1.1, ly - 0.3, 2.2, 0.5, h, size=17, bold=True, align="c")
            y = 1.9 + i * 1.35
            circle(s, 7.6, y + 0.45, 0.55, cols[i], str(i + 1), size=15)
            text(s, 8.1, y + 0.05, 4.6, 0.45, [[(h, {"bold": True}), (f"   {tag}", {"color": DIM, "size": 13})]], size=19)
            text(s, 8.1, y + 0.5, 4.6, 0.7, q, size=15, color=MUTED)
    with beat("zoom", gap=0.3, after=True, dur=0.6):
        text(s, 3.2, 3.75, 1.15, 0.5, "★", size=24, color=TEXT, align="c")
    with beat("rise", gap=0.2):
        text(s, 7.35, 6.0, 5.4, 0.5, sp["caption"], size=15, color=col, bold=True)
    return s


def s_process(m, sp):
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", MODULE_COLORS[m["n"]])
    title(s, sp["title"])
    cw, gap = 2.2, 0.28
    x0 = (W - (cw * 5 + gap * 4)) / 2
    # agrupadores
    for (a, b_, lab, c) in [(0, 1, "ENTENDER O PROBLEMA", LIME), (2, 4, "CRIAR E VALIDAR A SOLUÇÃO", TEAL)]:
        xa = x0 + a * (cw + gap)
        xb = x0 + b_ * (cw + gap) + cw
        with beat("wipe", gap=0.2):
            text(s, xa, 1.62, xb - xa, 0.35, [[(lab, {"cs": 200})]], size=11, color=c, bold=True, align="c")
            ln = A.track(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xa + 0.1), Inches(2.02), Inches(xb - 0.1), Inches(2.02)))
            ln.line.color.rgb = C(c)
            ln.line.width = Pt(1.5)
    for i, st in enumerate(STAGES):
        x = x0 + i * (cw + gap)
        c = STAGE_COLORS[st]
        with beat("rise", gap=0.3 if i else 0.25):
            box(s, x, 2.3, cw, 3.7, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, x, 3.9, cw, 0.5, st, size=19, bold=True, align="c")
            text(s, x + 0.15, 4.5, cw - 0.3, 1.2, STAGE_TXT[i], size=14, color=MUTED, align="c")
        with beat("zoom", gap=0.12):
            circle(s, x + cw / 2, 3.15, 1.0, c, str(i + 1), size=26)
        if i < 4:
            with beat("wipe", gap=0.25, dur=0.3):
                arrow_line(s, x + cw + 0.03, 3.15, x + cw + gap - 0.03, 3.15, color=DIM, w=2)
    with beat("fade", gap=0.4):
        text(s, 0.6, 6.3, 12.1, 0.4, "Módulos 2 a 5 aprofundam cada etapa. O Módulo 6 mostra tudo junto em um caso real.", size=14, color=DIM, align="c")
    return s


def s_loop(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    pw, gap = 2.1, 0.35
    x0 = (W - (pw * 5 + gap * 4)) / 2
    y = 3.6
    for i, st in enumerate(STAGES):
        x = x0 + i * (pw + gap)
        with beat("zoom", gap=0.15 if i else 0.25):
            p = box(s, x, y, pw, 0.7, fill=STAGE_COLORS[st], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            text(s, 0, 0, 0, 0, st, size=15, bold=True, color=BG, align="c", anchor="m", shape=p)
        if i < 4:
            with beat("wipe", gap=0.12, dur=0.25):
                arrow_line(s, x + pw + 0.03, y + 0.35, x + pw + gap - 0.03, y + 0.35, color=DIM, w=2)
    # arcos de retorno (iteração)
    for a, b_, hgt in [(4, 1, 1.5), (3, 2, 0.9), (4, 0, 2.1)]:
        xa = x0 + a * (pw + gap) + pw / 2
        xb = x0 + b_ * (pw + gap) + pw / 2
        with beat("wipe", gap=0.45, dur=0.7):
            arc = A.track(s.shapes.add_shape(MSO_SHAPE.ARC, Inches(xb), Inches(y - hgt), Inches(xa - xb), Inches(hgt * 2)))
        arc.adjustments[0] = 108.0   # 180° (unid. 60000/100000)
        arc.adjustments[1] = 0.0
        arc.fill.background()
        arc.shadow.inherit = False
        arc.line.color.rgb = C(MUTED)
        arc.line.width = Pt(1.5)
        arc.line.dash_style = MSO_LINE.DASH
        ln = arc.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:headEnd")).set("type", "triangle")
    with beat("fade", gap=0.3, after=True):
        text(s, 1.2, 4.75, 10.9, 1.5, sp["text"], size=17, color=MUTED, align="c")
    return s


def s_activity(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  Dinâmica", col)
    title(s, sp["title"])
    with beat("rise", gap=0.2):
        box(s, 0.6, 1.7, 3.9, 4.75, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
        text(s, 0.95, 1.95, 3.3, 0.4, [[("DURAÇÃO", {"cs": 200})]], size=11, color=BG, bold=True)
    with beat("zoom", gap=0.3, dur=0.55):
        text(s, 0.95, 2.25, 3.3, 1.0, sp["time"], size=44, color=BG, bold=True)
    with beat("fade", gap=0.25):
        text(s, 0.95, 3.4, 3.3, 0.4, [[("OBJETIVO", {"cs": 200})]], size=11, color=BG, bold=True)
        text(s, 0.95, 3.75, 3.3, 1.6, sp["goal"], size=15, color=BG)
        text(s, 0.95, 5.35, 3.3, 0.9, [[("Materiais: ", {"bold": True}), (sp["materials"], {})]], size=12, color=BG)
    n = len(sp["steps"])
    rowh = min(0.95, 4.75 / n)
    for i, st in enumerate(sp["steps"]):
        y = 1.7 + i * rowh
        with beat("rise", gap=0.2):
            circle(s, 5.15, y + rowh / 2, 0.5, SURFACE2, str(i + 1), size=14, color=col)
            text(s, 5.6, y, 7.1, rowh, st, size=16, anchor="m")
    return s


def s_csd(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    cols = [GREEN, LIME, BLUE]
    for i, (h, sub, notes) in enumerate(sp["cols"]):
        x = 0.6 + i * 4.1
        with beat("rise", gap=0.2):
            box(s, x, 1.7, 3.8, 4.8, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
            circle(s, x + 0.6, 2.25, 0.7, cols[i], h[0], size=22)
            text(s, x + 1.1, 1.93, 2.6, 0.5, h, size=21, bold=True)
            text(s, x + 1.1, 2.38, 2.6, 0.6, sub, size=12, color=MUTED)
    for i, (h, sub, notes) in enumerate(sp["cols"]):
        x = 0.6 + i * 4.1
        for j, n_ in enumerate(notes):
            py = 3.25 + j * 1.55
            with beat("zoom", gap=0.22 if (i or j) else 0.35, dur=0.4):
                pi = box(s, x + 0.35 + (j % 2) * 0.12, py, 3.0, 1.3, fill=cols[i], shape=MSO_SHAPE.RECTANGLE)
                pi.rotation = -1.5 if j % 2 == 0 else 1.2
                text(s, 0, 0, 0, 0, n_, size=14, color=BG, bold=True, align="c", anchor="m", shape=pi)
    return s


def s_steps(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    n = len(sp["steps"])
    cw = 12.13 / n
    with beat("wipe", gap=0.2, dur=0.3 * n):
        ln = A.track(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6 + cw / 2), Inches(2.85), Inches(0.6 + cw * (n - 0.5)), Inches(2.85)))
        ln.line.color.rgb = C(LINE)
        ln.line.width = Pt(3)
    for i, (h, b) in enumerate(sp["steps"]):
        cx = 0.6 + cw * (i + 0.5)
        with beat("zoom", gap=0.3 if i else 0.0):
            circle(s, cx, 2.85, 1.0, col, f"{i + 1:02d}", size=20)
        with beat("rise", gap=0.12):
            box(s, cx - cw / 2 + 0.12, 3.75, cw - 0.24, 2.6, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, cx - cw / 2 + 0.3, 3.95, cw - 0.6, 0.5, h, size=19, bold=True, align="c")
            text(s, cx - cw / 2 + 0.3, 4.5, cw - 0.6, 1.7, b, size=14, color=MUTED, align="c")
    return s


def s_affinity(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    import random
    rnd = random.Random(7)
    pcols = [LIME, GREEN, TEAL, BLUE]
    # caos
    with beat("fade", gap=0.2):
        box(s, 0.6, 1.7, 4.6, 4.2, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
        text(s, 0.6, 6.0, 4.6, 0.4, "Achados soltos", size=14, color=MUTED, align="c")
    with beat("zoom", gap=0.2, dur=0.35, stagger=0.08):
        for k in range(12):
            x = 0.85 + rnd.random() * 3.6
            y = 1.95 + rnd.random() * 3.2
            p = box(s, x, y, 0.62, 0.62, fill=pcols[k % 4])
            p.rotation = rnd.uniform(-14, 14)
    with beat("wipe", gap=0.3, after=True, dur=0.5):
        arrow_line(s, 5.45, 3.8, 6.45, 3.8, color=col, w=3)
    # grupos
    for gi, (g, n) in enumerate(sp["groups"]):
        x = 6.75 + (gi % 2) * 3.05
        y = 1.7 + (gi // 2) * 2.15
        with beat("rise", gap=0.3):
            box(s, x, y, 2.85, 1.95, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, x + 0.2, y + 0.12, 2.5, 0.4, g, size=16, bold=True, color=pcols[gi])
            text(s, x + 0.2, y + 1.35, 2.5, 0.4, f"{n} achados", size=11, color=DIM)
        with beat("zoom", gap=0.15, dur=0.3, stagger=0.07):
            for k in range(n):
                box(s, x + 0.22 + k * 0.62, y + 0.7, 0.5, 0.5, fill=pcols[gi])
    with beat("words", gap=0.3, after=True):
        text(s, 6.75, 6.0, 5.95, 0.4, "Padrões = insights", size=14, color=MUTED, align="c")
    with beat("fade", gap=0.3):
        text(s, 0.6, 6.4, 12, 0.5, sp["text"], size=13, color=DIM, align="c")
    return s


def s_persona(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    with beat("rise", gap=0.2):
        box(s, 0.6, 1.7, 3.9, 4.8, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    initials = "".join(w[0] for w in sp["name"].split(",")[0].split()[:2])
    with beat("zoom", gap=0.2, dur=0.6):
        circle(s, 2.55, 3.0, 1.7, col, initials, size=40)
    with beat("rise", gap=0.25):
        text(s, 0.8, 4.1, 3.5, 0.6, sp["name"], size=22, bold=True, align="c")
        text(s, 0.8, 4.7, 3.5, 1.0, sp["role"], size=14, color=MUTED, align="c")
        text(s, 0.8, 5.85, 3.5, 0.4, [[("PERSONA FICTÍCIA", {"cs": 200})]], size=10, color=DIM, align="c", bold=True)
    for i, (h, b) in enumerate(sp["fields"]):
        x = 4.8 + (i % 2) * 4.0
        y = 1.7 + (i // 2) * 2.45
        with beat("rise", gap=0.2 if i else 0.35):
            box(s, x, y, 3.8, 2.3, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, x + 0.3, y + 0.25, 3.2, 0.4, [[(h.upper(), {"cs": 150})]], size=11, color=col, bold=True)
            text(s, x + 0.3, y + 0.7, 3.2, 1.5, b, size=16, italic=h == "Frase")
    return s


def s_pyramid(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    levels = sp["levels"]
    n = len(levels)
    # Laterais contínuas até a ponta: o topo é um triângulo 1.7x mais alto que os demais níveis.
    base_w, apex_y, lh = 6.6, 1.55, 0.84
    cx = 3.9
    cols = [DEEP, BLUE, TEAL, GREEN, LIME]
    for i in reversed(range(n)):
        lab = levels[i]
        wb = base_w * (i + 1.7) / (n + 0.7)
        wt = base_w * (i + 0.7) / (n + 0.7)
        with beat("rise", gap=0.3 if i < n - 1 else 0.25):
            if i == 0:
                y, hh = apex_y, 1.7 * lh - 0.06
                box(s, cx - wb / 2, y, wb, hh, fill=cols[i], shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
                text(s, cx - 1.0, y + hh - 0.42, 2.0, 0.36, lab, size=11, bold=True, color=BG, align="c", anchor="m")
                continue
            y, hh = apex_y + (i + 0.7) * lh, lh - 0.06
            shp = box(s, cx - wb / 2, y, wb, hh, fill=cols[i], shape=MSO_SHAPE.TRAPEZOID)
            shp.adjustments[0] = ((wb - wt) / 2) / min(wb, hh)
            text(s, cx - 1.6, y, 3.2, hh, lab, size=15, bold=True, color=BG, align="c", anchor="m")
    with beat("fade", gap=0.4):
        text(s, 7.5, 2.0, 5.2, 3.5, sp["text"], size=18, color=MUTED, line_spacing=1.15)
    with beat("rise", gap=0.4):
        text(s, 7.5, 5.2, 5.2, 0.8, [[("Base primeiro: ", {"bold": True, "color": col}), ("atenda o essencial antes de encantar.", {})]], size=16)
    return s


def s_jobstory(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    labels = ["SITUAÇÃO", "MOTIVAÇÃO", "RESULTADO"]
    cols = [LIME, TEAL, BLUE]
    for i, (k, v) in enumerate(sp["parts"]):
        y = 1.75 + i * 1.5
        with beat("rise", gap=0.45 if i else 0.25):
            box(s, 0.6, y, 12.13, 1.3, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
            text(s, 0.95, y + 0.12, 2.8, 0.4, [[(labels[i], {"cs": 200})]], size=10, color=DIM, bold=True)
            text(s, 0.95, y + 0.42, 2.8, 0.7, k, size=26, bold=True, color=cols[i])
        with beat("words", gap=0.2, words_pct=25):
            text(s, 3.9, y, 8.6, 1.3, v, size=22, anchor="m")
    with beat("fade", gap=0.3, after=True):
        text(s, 0.6, 6.3, 12.1, 0.4, "Foque na situação e na motivação, não no perfil demográfico.", size=14, color=DIM)
    return s


def s_empathy(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    x0, y0, qw, qh = 0.6, 1.6, 6.0, 1.75
    cols = [LIME, GREEN, TEAL, BLUE]
    for i, (h, b) in enumerate(sp["quads"]):
        x = x0 + (i % 2) * (qw + 0.13)
        y = y0 + (i // 2) * (qh + 0.13)
        with beat("rise", gap=0.2):
            box(s, x, y, qw, qh, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            al = "l" if i % 2 == 0 else "r"
            tx = x + 0.35 if al == "l" else x + qw - 3.6
            text(s, tx, y + 0.3, 3.25, 0.5, h, size=20, bold=True, color=cols[i], align=al)
            text(s, tx, y + 0.85, 3.25, 0.8, b, size=14, color=MUTED, align=al)
    with beat("zoom", gap=0.35, dur=0.6):
        circle(s, x0 + qw + 0.065, y0 + qh + 0.065, 1.5, BG, None)
        circle(s, x0 + qw + 0.065, y0 + qh + 0.065, 1.3, col, "Persona", size=14)
    for i, (h, b) in enumerate(sp["bottom"]):
        x = x0 + i * (qw + 0.13)
        y = y0 + 2 * (qh + 0.13) + 0.05
        with beat("rise", gap=0.3):
            box(s, x, y, qw, 1.25, fill=SURFACE2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
            text(s, x + 0.35, y + 0.15, 5.3, 0.5, h, size=19, bold=True, color=[LIME, TEAL][i])
            text(s, x + 0.35, y + 0.62, 5.3, 0.5, b, size=14, color=MUTED)
    return s


def s_formula(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    with beat("rise", gap=0.2):
        box(s, 0.6, 1.8, 12.13, 1.8, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    runs = []
    for part, c in [("[Persona]", LIME), (" precisa de ", TEXT), ("[necessidade]", TEAL), (" porque ", TEXT), ("[insight]", BLUE), (".", TEXT)]:
        runs.append((part, {"color": c, "bold": part.startswith("[")}))
    with beat("words", gap=0.35, dur=0.5, words_pct=60):
        text(s, 0.9, 1.8, 11.5, 1.8, [runs], size=34, align="c", anchor="m")
    with beat("fade", gap=0.4, after=True):
        text(s, 0.6, 4.0, 3, 0.4, [[("EXEMPLO", {"cs": 200})]], size=11, color=col, bold=True)
    with beat("rise", gap=0.2):
        text(s, 0.6, 4.4, 12.1, 1.8, sp["example"], size=20, color=MUTED, italic=True, line_spacing=1.1)
    return s


def s_hmw(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    with beat("words", dur=0.5, words_pct=70):
        text(s, 0.6, 1.1, 5.0, 2.8, [[("Como", {})], [("poderíamos", {})], [("…?", {"color": col})]], size=50, bold=True, line_spacing=0.95)
    with beat("fade", gap=0.2, after=True):
        text(s, 0.6, 4.35, 4.6, 1.5, sp["text"], size=17, color=MUTED)
    for i, ex in enumerate(sp["examples"]):
        y = 1.3 + i * 1.7
        with beat("rise", gap=0.35):
            box(s, 5.9, y, 6.83, 1.45, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
            text(s, 6.95, y, 5.6, 1.45, ex, size=17, anchor="m")
        with beat("zoom", gap=0.12):
            circle(s, 6.45, y + 0.72, 0.5, col, "?", size=16)
    return s


def s_diverge(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    for i, (lab, c, sub) in enumerate([("Problema", LIME, "Empatizar  →  Definir"), ("Solução", TEAL, "Idear  →  Prototipar e Testar")]):
        x = 0.9 + i * 5.9
        with beat("wipe", gap=0.6 if i else 0.25, dur=1.0):
            d = box(s, x, 1.9, 5.6, 3.1, fill=c, shape=MSO_SHAPE.DIAMOND, alpha=30)
            d.line.color.rgb = C(c)
            d.line.width = Pt(2)
        with beat("fade", gap=0.1):
            text(s, x + 0.2, 5.1, 2.6, 0.4, "divergir", size=14, color=c, bold=True, align="c")
        with beat("fade", gap=0.5):
            text(s, x + 2.8, 5.1, 2.6, 0.4, "convergir", size=14, color=c, bold=True, align="c")
        with beat("zoom", gap=0.0):
            text(s, x, 1.9, 5.6, 3.1, [[(lab, {"bold": True, "size": 24})], [(sub, {"size": 13, "color": MUTED})]], size=20, align="c", anchor="m")
    with beat("fade", gap=0.4, after=True):
        text(s, 0.6, 5.75, 12.1, 0.8, sp["text"], size=17, color=MUTED, align="c")
    return s


def s_crazy8(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    with beat("fade", gap=0.2):
        text(s, 0.6, 1.8, 4.3, 3.0, sp["text"], size=18, color=MUTED, line_spacing=1.15)
    gx, gy, cw, ch = 5.4, 1.75, 1.75, 2.25
    with beat("rise", gap=0.3):
        box(s, gx - 0.1, gy - 0.1, cw * 4 + 0.2, ch * 2 + 0.2, fill=TEXT, shape=MSO_SHAPE.RECTANGLE)
    for i in range(8):
        x = gx + (i % 4) * cw
        y = gy + (i // 4) * ch
        with beat("zoom", gap=0.2 if i else 0.35, dur=0.35):
            box(s, x, y, cw, ch, fill="FFFFFF", line="C9CCC4", lw=1, dash=True)
            text(s, x + 0.1, y + 0.08, 0.6, 0.4, str(i + 1), size=14, bold=True, color=col if col != LIME else GREEN)
            text(s, x, y + ch - 0.45, cw - 0.1, 0.35, "1 min", size=10, color="8A8F87", align="r")
    with beat("zoom", gap=0.3, after=True, dur=0.6):
        text(s, 0.6, 4.6, 4.3, 1.4, [[("8", {"color": col})], [("ideias em 8 minutos", {"size": 16, "color": TEXT})]], size=64, bold=True, line_spacing=0.9)
    return s


def s_matrix(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    x0, y0, qw, qh = 1.6, 1.7, 5.4, 2.25
    with beat("wipe_up", gap=0.2):
        arrow_line(s, 1.25, y0 + qh * 2 + 0.12, 1.25, y0, color=MUTED, w=2)
        text(s, 0.35, 3.3, 0.8, 1.6, "IMPACTO", size=11, color=MUTED, bold=True, align="c", anchor="m").rotation = -90
    with beat("wipe", gap=0.2):
        arrow_line(s, x0, 6.55, x0 + qw * 2 + 0.12, 6.55, color=MUTED, w=2)
        text(s, x0, 6.6, qw * 2, 0.3, "ESFORÇO  →", size=11, color=MUTED, bold=True, align="c")
    order = [(1, 1, 3, "C0584F"), (0, 1, 2, DIM), (1, 0, 1, BLUE), (0, 0, 0, LIME)]
    for cx_, cy_, k, c in order:
        h, b = sp["quads"][k]
        x = x0 + cx_ * (qw + 0.12)
        y = y0 + cy_ * (qh + 0.12)
        with beat("rise", gap=0.3):
            box(s, x, y, qw, qh, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
            text(s, x + 0.9, y + 0.3, qw - 1.1, 0.5, h, size=20, bold=True, color=c if c != DIM else MUTED)
            text(s, x + 0.9, y + 0.85, qw - 1.1, 1.2, b, size=15, color=MUTED)
        with beat("zoom", gap=0.12, dur=0.6 if k == 0 else 0.45):
            circle(s, x + 0.5, y + 0.55, 0.42, c, None)
    return s


def s_grid4(m, sp):
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"])
    chrome(s, f"Módulo {m['n']}  ·  {m['title']}", col)
    title(s, sp["title"])
    icons = ["+", "−", "?", "★"]
    cols = [LIME, "C0584F", BLUE, TEAL]
    for i, (h, b) in enumerate(sp["quads"]):
        x = 0.6 + (i % 2) * 6.15
        y = 1.7 + (i // 2) * 2.45
        with beat("rise", gap=0.25):
            box(s, x, y, 5.95, 2.25, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, x + 1.7, y + 0.55, 4.0, 0.6, h, size=22, bold=True)
            text(s, x + 1.7, y + 1.15, 4.0, 0.8, b, size=16, color=MUTED)
        with beat("zoom", gap=0.15):
            circle(s, x + 0.85, y + 1.12, 1.0, cols[i], icons[i], size=30)
    return s


def s_impact(m, sp):
    """Frase de impacto: entra palavra por palavra, destaques na cor do módulo, número gigante ao fundo."""
    col = MODULE_COLORS[m["n"]]
    s = new_slide(sp["notes"], transition="black")
    with beat("fade", dur=1.4):
        text(s, 6.6, -0.35, 6.4, 6.0, f"{m['n']:02d}", size=330, bold=True, color=SURFACE, align="r", font=FONT_B)
    with beat("wipe_up", gap=0.3, dur=0.7):
        box(s, 0.9, 2.55, 0.09, 1.55, fill=col)
    with beat("fade", gap=0.35):
        text(s, 1.3, 1.95, 9, 0.4, [[(f"PARA LEVAR  ·  MÓDULO {m['n']}", {"cs": 300})]], size=12, color=col, bold=True)
    paras = []
    for line in sp["lines"]:
        parts = line.split("*")
        paras.append([(t, {"color": col} if k % 2 else {}) for k, t in enumerate(parts) if t])
    with beat("words", gap=0.3, dur=0.5, words_pct=45):
        text(s, 1.3, 2.45, 11.2, 1.8, paras, size=48, bold=True, line_spacing=1.05)
    with beat("wipe", gap=0.4, after=True, dur=0.8):
        box(s, 1.3, 4.55, 2.6, 0.06, fill=col)
    with beat("rise", gap=0.3):
        text(s, 1.3, 4.8, 10.8, 0.8, sp["sub"], size=18, color=MUTED)
    text(s, 0.6, 7.0, 8, 0.3, f"{B.BRAND}  ·  {B.COURSE}", size=10, color=DIM)
    return s


# ---- estudo de caso ----
def case_chrome(s, sp, col):
    chrome(s, "Módulo 6  ·  Estudo de caso (ilustrativo)", col)


def s_case_intro(m, sp):
    col = MODULE_COLORS[6]
    s = new_slide(sp["notes"])
    case_chrome(s, sp, col)
    title(s, sp["title"], w=8)
    with beat("fade", gap=0.25):
        text(s, 0.6, 1.8, 7.4, 3.5, sp["text"], size=19, color=MUTED, line_spacing=1.15)
    with beat("rise", gap=0.4):
        box(s, 8.6, 1.7, 4.13, 4.3, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        text(s, 8.9, 4.0, 3.5, 0.8, sp["stat"][1], size=18, color=MUTED, align="c")
    with beat("zoom", gap=0.2, dur=0.65):
        text(s, 8.6, 2.3, 4.13, 1.6, sp["stat"][0], size=88, bold=True, color=LIME, align="c")
    # a pergunta só aparece no clique: dê tempo para a turma pensar no problema
    with beat("words", click=True, words_pct=30):
        text(s, 0.6, 5.5, 7.4, 0.6, "Qual seria a primeira solução que você proporia?", size=18, bold=True, color=col)
    return s


def s_case_stage(m, sp):
    col = STAGE_COLORS[sp["stage"]]
    s = new_slide(sp["notes"])
    case_chrome(s, sp, col)
    title(s, sp["title"])
    n = len(sp["bullets"])
    gap = 0.14
    rh = (4.35 - gap * (n - 1)) / n
    for i, b_ in enumerate(sp["bullets"]):
        lab, _, body = b_.partition(": ")
        y = 1.65 + i * (rh + gap)
        with beat("rise", gap=0.3 if i else 0.25):
            box(s, 0.6, y, 12.13, rh, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
            text(s, 0.9, y, 2.9, rh, lab, size=15, bold=True, color=col, anchor="m")
            text(s, 3.85, y, 8.7, rh, body, size=16, anchor="m")
    with beat("wipe", gap=0.35, dur=0.8):
        stage_bar(s, sp["stage"], x=0.6, y=6.3, w=12.13)
    return s


def s_case_persona(m, sp):
    col = MODULE_COLORS[6]
    s = new_slide(sp["notes"])
    case_chrome(s, sp, col)
    title(s, sp["title"])
    for i, (nm, role, need) in enumerate(sp["personas"]):
        x = 0.6 + i * 6.15
        c = [LIME, TEAL][i]
        with beat("rise", gap=0.5 if i else 0.25):
            box(s, x, 1.7, 5.95, 4.7, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
            text(s, x + 2.0, 2.3, 3.7, 0.6, nm, size=24, bold=True)
            text(s, x + 2.0, 2.9, 3.7, 0.8, role, size=14, color=MUTED)
        ini = "".join(w[0] for w in nm.split(",")[0].split()[:2])
        with beat("zoom", gap=0.15, dur=0.55):
            circle(s, x + 1.1, 2.75, 1.35, c, ini, size=30)
        with beat("fade", gap=0.3):
            text(s, x + 0.45, 4.0, 5, 0.4, [[("NECESSIDADE", {"cs": 150})]], size=11, color=c, bold=True)
            text(s, x + 0.45, 4.4, 5.1, 1.8, need, size=19)
    return s


def s_case_ideas(m, sp):
    col = TEAL
    s = new_slide(sp["notes"])
    case_chrome(s, sp, col)
    title(s, sp["title"])
    groups = [("win", "Vitórias rápidas", LIME, "Prototipar já"), ("bet", "Grandes apostas", BLUE, "Depois do piloto"), ("avoid", "Evitar", "C0584F", "Descartadas")]
    for gi, (k, lab, c, sub) in enumerate(groups):
        x = 0.6 + gi * 4.1
        with beat("rise", gap=0.2):
            box(s, x, 1.7, 3.8, 4.7, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
            text(s, x + 0.3, 1.9, 3.2, 0.5, lab, size=20, bold=True, color=c)
            text(s, x + 0.3, 2.4, 3.2, 0.4, sub, size=12, color=DIM)
    for gi, (k, lab, c, sub) in enumerate(groups):
        x = 0.6 + gi * 4.1
        yy = 3.0
        for idea, kk in sp["ideas"]:
            if kk != k:
                continue
            with beat("zoom", gap=0.25):
                p = box(s, x + 0.3, yy, 3.2, 1.2, fill=SURFACE2, line=c, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
                text(s, 0, 0, 0, 0, idea, size=15, bold=True, align="c", anchor="m", shape=p)
            yy += 1.4
    return s


def s_case_result(m, sp):
    col = DEEP
    s = new_slide(sp["notes"])
    case_chrome(s, sp, BLUE)
    title(s, sp["title"])
    cols = [LIME, TEAL, BLUE]
    for i, (num, lab) in enumerate(sp["stats"]):
        x = 0.6 + i * 4.1
        with beat("rise", gap=0.35 if i else 0.25):
            box(s, x, 1.75, 3.8, 2.8, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, x + 0.3, 3.4, 3.2, 1.0, lab, size=15, color=MUTED, align="c")
        with beat("zoom", gap=0.15, dur=0.55):
            text(s, x, 2.1, 3.8, 1.2, num, size=46, bold=True, color=cols[i], align="c")
    with beat("fade", gap=0.4):
        text(s, 0.6, 4.9, 12.1, 1.0, sp["text"], size=17, color=MUTED)
        text(s, 0.6, 6.05, 12.1, 0.4, "Números ilustrativos, para fins didáticos.", size=11, color=DIM, italic=True)
    return s


RENDER = {
    "quote": s_quote, "cards": s_cards, "split": s_split, "bullets": s_bullets, "stat": s_stat,
    "pillars": s_pillars, "process": s_process, "loop": s_loop, "activity": s_activity, "csd": s_csd,
    "steps": s_steps, "affinity": s_affinity, "persona": s_persona, "pyramid": s_pyramid,
    "jobstory": s_jobstory, "empathy": s_empathy, "formula": s_formula, "hmw": s_hmw,
    "diverge": s_diverge, "crazy8": s_crazy8, "matrix": s_matrix, "grid4": s_grid4,
    "case_intro": s_case_intro, "case_stage": s_case_stage, "case_persona": s_case_persona,
    "case_ideas": s_case_ideas, "case_result": s_case_result, "impact": s_impact,
}


# ---- desafio, prova, encerramento ----
def s_divider(label, big, sub, col, notes):
    s = new_slide(notes, transition="black")
    with beat("fade"):
        text(s, 0.6, 2.2, 8, 0.4, [[(label, {"cs": 250})]], size=13, color=col, bold=True)
    with beat("words", gap=0.25, dur=0.5, words_pct=60):
        text(s, 0.6, 2.7, 11, 1.4, big, size=54, bold=True)
    with beat("rise", gap=0.3, after=True):
        text(s, 0.6, 4.05, 10, 1.0, sub, size=20, color=MUTED)
    circuit(s, 10.3, 0.7, color=col, scale=1.1)
    text(s, 0.6, 7.0, 8, 0.3, f"{B.BRAND}  ·  {B.COURSE}", size=10, color=DIM)
    return s


def s_challenge_brief():
    s = new_slide("Leia o enunciado com a turma. Confirme que cada equipe tem seu problema (o mesmo da Matriz CSD do Encontro 1). Entregue o canvas impresso.")
    chrome(s, "Desafio Final", LIME)
    title(s, "O desafio")
    with beat("fade", gap=0.2):
        text(s, 0.6, 1.6, 7.3, 3.0, CHALLENGE["brief"], size=19, color=MUTED, line_spacing=1.15)
    with beat("rise", gap=0.35):
        box(s, 8.4, 1.7, 4.33, 3.2, fill=LIME, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        text(s, 8.7, 1.95, 3.8, 0.4, [[("TEMPO", {"cs": 200})]], size=11, color=BG, bold=True)
        text(s, 8.7, 3.45, 3.8, 1.2, "+ pitch de 5 minutos por equipe", size=17, color=BG)
    with beat("zoom", gap=0.2, dur=0.6):
        text(s, 8.7, 2.3, 3.8, 1.0, "90 min", size=52, bold=True, color=BG)
    with beat("fade", gap=0.35):
        text(s, 0.6, 5.2, 12, 0.4, [[("CRONOGRAMA", {"cs": 200})]], size=11, color=LIME, bold=True)
    tl = CHALLENGE["timeline"]
    cw = 12.13 / len(tl)
    for i, (t, d) in enumerate(tl):
        x = 0.6 + i * cw
        with beat("rise", gap=0.15):
            box(s, x + 0.04, 5.6, cw - 0.08, 1.1, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
            text(s, x + 0.12, 5.64, cw - 0.24, 0.35, t, size=11, bold=True, color=[LIME, GREEN, TEAL, BLUE, DEEP, MUTED][i])
            text(s, x + 0.12, 5.95, cw - 0.24, 0.72, d, size=10.5, color=MUTED)
    return s


def s_challenge_deliv():
    s = new_slide("Estas são as entregas mínimas. Equipes que terminarem antes podem fazer um segundo ciclo de teste.")
    chrome(s, "Desafio Final", LIME)
    title(s, "O que cada equipe entrega")
    for i, (st, d) in enumerate(CHALLENGE["deliverables"]):
        y = 1.65 + i * 0.97
        c = STAGE_COLORS[st]
        with beat("zoom", gap=0.3 if i else 0.25):
            p = box(s, 0.6, y, 2.6, 0.8, fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
            text(s, 0, 0, 0, 0, st, size=17, bold=True, color=BG, align="c", anchor="m", shape=p)
        with beat("wipe", gap=0.15):
            box(s, 3.35, y, 9.38, 0.8, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
            text(s, 3.7, y, 8.9, 0.8, d, size=17, anchor="m")
    return s


def s_challenge_rubric():
    s = new_slide("Mostre a rubrica ANTES de começar: as equipes precisam saber como serão avaliadas. Use o gabarito do roteiro para pontuar cada pitch.")
    chrome(s, "Desafio Final", LIME)
    title(s, "Como avaliaremos o pitch")
    for i, (h, d, pts) in enumerate(CHALLENGE["rubric"]):
        x = 0.6 + (i % 2) * 6.15
        y = 1.7 + (i // 2) * 2.4
        with beat("rise", gap=0.25):
            box(s, x, y, 5.95, 2.2, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, x + 0.35, y + 0.3, 4, 0.5, h, size=20, bold=True)
            text(s, x + 0.35, y + 0.85, 4.2, 1.2, d, size=14, color=MUTED)
            text(s, x + 4.4, y + 1.1, 1.3, 0.3, "pontos", size=11, color=DIM, align="r")
        with beat("zoom", gap=0.15):
            text(s, x + 4.4, y + 0.3, 1.3, 0.8, f"{pts}", size=40, bold=True, color=[LIME, GREEN, TEAL, BLUE][i], align="r")
    with beat("zoom", gap=0.4, after=True, dur=0.6):
        text(s, 0.6, 6.5, 12, 0.3, "Total: 100 pontos", size=14, color=LIME, bold=True)
    return s


def s_exam_intro():
    s = new_slide("Distribua a prova impressa. 25 minutos, individual e sem consulta. Recolha antes de mostrar a correção comentada.")
    chrome(s, "Avaliação", BLUE)
    title(s, "Prova final")
    info = [("6", "questões de múltipla escolha"), ("25", "minutos"), ("1", "questão por módulo")]
    for i, (n, l) in enumerate(info):
        x = 0.6 + i * 4.1
        with beat("rise", gap=0.3 if i else 0.25):
            box(s, x, 1.9, 3.8, 2.6, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
            text(s, x + 0.3, 3.55, 3.2, 0.8, l, size=17, color=MUTED, align="c")
        with beat("zoom", gap=0.15, dur=0.55):
            text(s, x, 2.15, 3.8, 1.3, n, size=72, bold=True, color=[LIME, TEAL, BLUE][i], align="c")
    with beat("fade", gap=0.4):
        bullets(s, 0.6, 4.9, 12, 1.5, ["Individual e sem consulta", "Leia a situação com atenção: as questões são práticas", "Marque apenas uma alternativa por questão"], size=16, dot=BLUE, spacing=8)
    return s


def s_exam_answer(i, q):
    s = new_slide("Correção comentada. Pergunte à turma qual alternativa marcaram ANTES de revelar. "
                  "CLIQUE para revelar a resposta: a correta acende em verde e as erradas apagam. Depois leia a justificativa.\n\n" + q["why"])
    chrome(s, f"Correção comentada  ·  Questão {i + 1}", BLUE)
    with beat("fade"):
        text(s, 0.6, 0.65, 12.1, 1.6, q["q"], size=17, bold=True, line_spacing=1.08)
    for j, o in enumerate(q["opts"]):
        y = 2.45 + j * 0.78
        with beat("rise", gap=0.15 if j else 0.3):
            box(s, 0.6, y, 12.13, 0.66, fill=SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
            circle(s, 1.0, y + 0.33, 0.42, LINE, "abcd"[j], size=13, color=MUTED)
            text(s, 1.4, y, 11.1, 0.66, o, size=14, anchor="m", color=TEXT)
    # revelação (clique): véu escuro sobre as erradas + destaque na correta
    with beat("fade", click=True, dur=0.6):
        for j in range(len(q["opts"])):
            if j != q["ans"]:
                box(s, 0.55, 2.42 + j * 0.78, 12.23, 0.72, fill=BG, alpha=62)
    ya = 2.45 + q["ans"] * 0.78
    with beat("zoom", gap=0.1):
        box(s, 0.6, ya, 12.13, 0.66, line=LIME, lw=2.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
        circle(s, 1.0, ya + 0.33, 0.42, LIME, "abcd"[q["ans"]], size=13, color=BG)
    with beat("zoom", gap=0.25):
        b = box(s, 11.0, ya + 0.13, 1.55, 0.4, fill=LIME, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        text(s, 0, 0, 0, 0, "✓  Correta", size=12, color=BG, bold=True, align="c", anchor="m", shape=b)
    with beat("rise", gap=0.35):
        text(s, 0.6, 5.65, 12.1, 1.2, [[("Por quê: ", {"bold": True, "color": LIME}), (q["why"], {})]], size=12.5, color=MUTED)
    return s


def s_closing():
    s = new_slide("Agradeça, peça a avaliação do treinamento (formulário ou QR code) e ofereça os próximos passos: consultoria para aplicar o método ou projetos de dados.")
    A.transition("black")
    with beat("fade"):
        text(s, 0.8, 1.6, 6.3, 0.4, [[("OBRIGADO!", {"cs": 300})]], size=14, color=LIME, bold=True)
    with beat("words", gap=0.25, dur=0.5, words_pct=55):
        text(s, 0.8, 2.1, 6.3, 2.0, "Agora é colocar em prática.", size=44, bold=True, line_spacing=0.95)
    with beat("zoom", gap=0.2, dur=0.9):
        A.track(s.shapes.add_picture(str(ASSETS / "logo_full_t.png"), Inches(7.4), Inches(1.4), height=Inches(3.4)))
    with beat("rise", gap=0.4, after=True):
        text(s, 0.8, 4.0, 6.3, 1.0, "Comece pequeno: escolha um problema, converse com 5 clientes esta semana.", size=18, color=MUTED)
    with beat("rise", gap=0.3):
        text(s, 0.8, 5.3, 6.3, 1.0, [[(B.INSTRUCTOR, {"bold": True})], [("hayaidatasystems@gmail.com", {"color": LIME})]], size=16, spacing=4)
    text(s, 0.8, 6.85, 8, 0.3, f"{B.BRAND}  ·  {B.TAGLINE}", size=11, color=DIM)
    return s


# ---------------------------------------------------------------- montagem
def build(out):
    cover = s_cover()
    welcome = s_welcome()
    with beat("zoom", gap=0.3, after=True):
        btn = button(welcome, 0.6, 5.55, 3.4, 0.55, "▶  Vídeo de boas-vindas", fill=LIME)
    vid = s_video("abertura", "Boas-vindas", welcome, LIME)
    btn.click_action.target_slide = vid
    s_about()
    s_rules()
    s_agenda()
    for m in MODULES:
        sec, btn = s_section(m)
        vid = s_video(m["video"], f"Módulo {m['n']}", sec, MODULE_COLORS[m["n"]])
        btn.click_action.target_slide = vid
        for sp in m["slides"]:
            RENDER[sp["type"]](m, sp)
    s_divider("ENCONTRO 3  ·  DESAFIO FINAL", "Mão na massa", "Aplique as 5 etapas em um problema real da sua empresa.", LIME,
              "Transição para o Desafio Final. Organize as equipes nas mesas com os materiais.")
    s_challenge_brief()
    s_challenge_deliv()
    s_challenge_rubric()
    s_divider("ENCONTRO 3  ·  AVALIAÇÃO", "Hora da prova", "6 questões para consolidar o que aprendemos.", BLUE,
              "Transição para a prova.")
    s_exam_intro()
    for i, q in enumerate(EXAM):
        s_exam_answer(i, q)
    s_closing()
    A.finish()
    prs.save(out)
    return len(prs.slides)


if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else str(B.ROOT / "DT_na_Pratica_Slides.pptx")
    n = build(out)
    print("slides:", n, "->", out)
